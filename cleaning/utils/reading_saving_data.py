from cleaning.utils.utils_webapp import create_colnames_dict
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font
from openpyxl import Workbook
import pandas as pd


def read_exports(LPO_file_list):
    exports_list = list()
    for file_path in LPO_file_list:
        if file_path.endswith(".csv"):
            exports_list.append(pd.read_csv(file_path))
        elif file_path.endswith(".xlsx"):
            exports_list.append(pd.read_excel(file_path))
    return exports_list

def style_and_save(cleaned_data, raw_data, saving_file_path):
    wb = Workbook()
    
    #### Filling the cleaned sheet
    ws = wb.active
    ws.title = "Cleaned Export LPO"
    for row in dataframe_to_rows(cleaned_data, 
                                    header=True, index=False):
        ws.append(row)
    style = TableStyleInfo(name="TableStyleMedium8", showRowStripes=True)

    table = Table(displayName="cleaned_export_LPO", 
              ref="A1:" + get_column_letter(ws.max_column) + str(ws.max_row))
    table.tableStyleInfo = style
    ws.add_table(table)

    # Setting Columns'width
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = 25
    
    # Highlighting aberrant value in CD_NOM (negative or NULL)
    colnames_index = create_colnames_dict(worksheet=ws)
    for row_cells in ws.iter_rows(min_row=ws.min_row + 1 , max_row=ws.max_row):
        for cell in row_cells:
            if pd.isnull(row_cells[colnames_index['CD_NOM']].value) or (row_cells[colnames_index['CD_NOM']].value < 0):
                cell.fill = PatternFill(fill_type='solid', start_color='d49137', end_color='d49137')
                cell.font = Font(color="FFFFFF")
    
    #### Creating and filling the raw sheet
    wb.create_sheet('Raw Export LPO')
    wb.active = wb['Raw Export LPO']
    ws_raw = wb.active
    for row in dataframe_to_rows(raw_data, 
                                    header=True, index=False):
        ws_raw.append(row)
    style = TableStyleInfo(name="TableStyleLight15", showRowStripes=True)

    table = Table(displayName="raw_export_LPO", 
              ref="A1:" + get_column_letter(ws_raw.max_column) + str(ws_raw.max_row))
    table.tableStyleInfo = style
    ws_raw.add_table(table)

    # Setting Columns'width
    for col in ws_raw.columns:
        ws_raw.column_dimensions[get_column_letter(col[0].column)].width = 25

    wb.active = wb['Cleaned Export LPO']
    wb.save(saving_file_path)


